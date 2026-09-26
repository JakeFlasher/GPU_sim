"""Export selected Nsight Compute report metrics to an XLSX workbook.

This utility is part of the local-label evidence workflow. It consumes existing
``.ncu-rep`` files; it does not launch NCU or capture new labels.
"""
from __future__ import annotations

import argparse
import glob
import importlib
import importlib.util
import os
import shutil
import sys
from pathlib import Path
from typing import Any, Iterable, Sequence

try:
    from ncu_report_helper import *  # type: ignore[import-not-found]  # noqa: F403
except ImportError:  # pragma: no cover - package-import path
    from tools.local_labels.ncu_report_helper import *  # type: ignore[F403]


_DEFAULT_OUTPUT = Path("ncu_report.xlsx")
_REPORT_PATH_ENVS = (
    "NCU_REPORT_PYTHON_PATH",
    "NSIGHT_COMPUTE_PYTHON_PATH",
)


def _iter_existing_paths(paths: Iterable[Path]) -> Iterable[Path]:
    seen: set[Path] = set()
    for path in paths:
        try:
            resolved = path.expanduser().resolve()
        except OSError:
            continue
        if resolved in seen or not resolved.exists():
            continue
        seen.add(resolved)
        yield resolved


def _find_ncu_report_module_dir(root: Path) -> Path | None:
    root = root.expanduser()
    if root.is_file() and root.name == "ncu_report.py":
        return root.parent.resolve()
    if not root.is_dir():
        return None
    direct = root / "ncu_report.py"
    if direct.is_file():
        return root.resolve()
    for dirpath, _dirnames, filenames in os.walk(root):
        if "ncu_report.py" in filenames:
            return Path(dirpath).resolve()
    return None


def find_nsight_folder(cuda_base_path: str | os.PathLike[str]) -> str | None:
    """Backward-compatible wrapper returning a directory with ncu_report.py."""
    found = _find_ncu_report_module_dir(Path(cuda_base_path))
    return str(found) if found is not None else None


def _candidate_search_roots() -> list[Path]:
    roots: list[Path] = []
    for env_name in _REPORT_PATH_ENVS:
        value = os.environ.get(env_name)
        if value:
            roots.extend(Path(part) for part in value.split(os.pathsep) if part)

    spec = importlib.util.find_spec("ncu_report")
    if spec is not None and spec.origin:
        roots.append(Path(spec.origin))

    ncu = shutil.which("ncu")
    if ncu:
        ncu_path = Path(ncu)
        roots.extend((ncu_path.parent, ncu_path.parent.parent))

    roots.append(Path("/usr/local/cuda"))
    roots.extend(Path(p) for p in sorted(glob.glob("/usr/local/cuda-*"), reverse=True))
    roots.append(Path("/opt/nvidia/nsight-compute"))
    roots.extend(
        Path(p) for p in sorted(glob.glob("/opt/nvidia/nsight-compute-*"), reverse=True)
    )
    return list(_iter_existing_paths(roots))


def discover_ncu_report_path(
    explicit_path: str | os.PathLike[str] | None = None,
    *,
    search_roots: Sequence[Path] | None = None,
) -> Path:
    """Return the directory containing Nsight Compute's ``ncu_report.py``."""
    candidates = []
    if explicit_path is not None:
        candidates.append(Path(explicit_path))
    candidates.extend(search_roots if search_roots is not None else _candidate_search_roots())
    for candidate in _iter_existing_paths(candidates):
        found = _find_ncu_report_module_dir(candidate)
        if found is not None:
            return found
    raise FileNotFoundError(
        "could not find Nsight Compute ncu_report.py; pass --ncu-report-path "
        "or set NCU_REPORT_PYTHON_PATH"
    )


def load_ncu_report_module(
    explicit_path: str | os.PathLike[str] | None = None,
    *,
    search_roots: Sequence[Path] | None = None,
) -> Any:
    module_dir = discover_ncu_report_path(
        explicit_path, search_roots=search_roots
    )
    if str(module_dir) not in sys.path:
        sys.path.insert(0, str(module_dir))
    return importlib.import_module("ncu_report")


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "reports",
        nargs="*",
        type=Path,
        help="Input .ncu-rep file paths.",
    )
    parser.add_argument(
        "--report",
        action="append",
        default=[],
        type=Path,
        help="Additional input .ncu-rep file path. May be repeated.",
    )
    parser.add_argument(
        "--report-glob",
        action="append",
        default=[],
        help="Glob for input .ncu-rep files. May be repeated.",
    )
    parser.add_argument(
        "--output",
        "-o",
        type=Path,
        default=_DEFAULT_OUTPUT,
        help=f"Output workbook path (default: {_DEFAULT_OUTPUT}).",
    )
    parser.add_argument(
        "--ncu-report-path",
        type=Path,
        help="Directory containing ncu_report.py, or the ncu_report.py file itself.",
    )
    parser.add_argument(
        "--max-kernels",
        type=int,
        default=100,
        help="Maximum kernels to export per report.",
    )
    return parser.parse_args(argv)


def _resolve_report_paths(args: argparse.Namespace) -> list[Path]:
    paths = [Path(p) for p in args.reports]
    paths.extend(Path(p) for p in args.report)
    for pattern in args.report_glob:
        paths.extend(Path(p) for p in sorted(glob.glob(pattern)))
    resolved = list(_iter_existing_paths(paths))
    if not resolved:
        raise FileNotFoundError(
            "no input .ncu-rep files found; pass positional report paths, "
            "--report, or --report-glob"
        )
    return resolved


def _header_row() -> list[str]:
    return [
        "",
        "Kernel ID",
        "Thread block Limit SM",
        "Thread block limit registers",
        "Thread block limit shared memory",
        "Thread block limit warps",
        "theoretical max active warps per SM",
        "theoretical occupancy",
        "achieved active warps per SM",
        "achieved occupancy",
        "unified L1 cache hit rate",
        "unified L1 cache hit rate for read transactions (global memory accesses)",
        "L2 cache hit rate",
        "GMEM read requests",
        "GMEM write requests",
        "GMEM total requests",
        "GMEM read transactions",
        "GMEM write transactions",
        "GMEM total transactions",
        "number of read transactions per read requests",
        "number of write transactions per write requests",
        "L2 read transactions",
        "L2 write transactions",
        "L2 total transactions",
        "DRAM total transactions",
        "Total number of global atomic requests",
        "Total number of global reduction requests",
        "Global memory atomic and reduction transactions",
        "GPU active cycles",
        "SM active cycles",
        "Warp instructions executed",
        "Instructions executed per clock cycle (IPC)",
        "Total instructions executed per seconds (MIPS)",
        "Kernel execution time (ns)",
        "unified L1 cache total requests",
        "unified L2 cache total requests",
        "DRAM bytes per second",
        "Memory busy (%)",
        "Max bandwidth (%)",
        "Memory pipes busy (%)",
        "DRAM clock (Hz)",
    ]


def _kernel_rows(
    *, report_file_path: Path, ncu_report: Any, max_kernels: int,
) -> list[list[Any]]:
    report = ncu_report.load_report(str(report_file_path))
    kernels = report[0]
    rows: list[list[Any]] = []
    for knum, kernel in enumerate(kernels[:max_kernels]):
        read_requests = get_l1tex__t_requests_pipe_lsu_mem_global_op_ld_sum(kernel)  # noqa: F405
        write_requests = get_l1tex__t_requests_pipe_lsu_mem_global_op_st_sum(kernel)  # noqa: F405
        read_transactions = get_l1tex__t_sectors_pipe_lsu_mem_global_op_ld_sum(kernel)  # noqa: F405
        write_transactions = get_l1tex__t_sectors_pipe_lsu_mem_global_op_st_sum(kernel)  # noqa: F405
        l2_read = get_l1tex__m_xbar2l1tex_read_sectors_mem_lg_op_ld_sum(kernel)  # noqa: F405
        l2_write = get_l1tex__m_l1tex2xbar_write_sectors_mem_lg_op_st_sum(kernel)  # noqa: F405
        atom_requests = get_l1tex__t_requests_pipe_lsu_mem_global_op_atom_sum(kernel)  # noqa: F405
        red_requests = get_l1tex__t_requests_pipe_lsu_mem_global_op_red_sum(kernel)  # noqa: F405
        ipc = get_sm__inst_executed_avg_per_cycle_elapsed(kernel)  # noqa: F405
        l1_total_requests = get_L1_Total_Requests(kernel)  # noqa: F405
        app_results: list[Any] = [
            report_file_path.stem,
            str(knum),
            get_launch__occupancy_limit_blocks(kernel),  # noqa: F405
            get_launch__occupancy_limit_registers(kernel),  # noqa: F405
            get_launch__occupancy_limit_shared_mem(kernel),  # noqa: F405
            get_launch__occupancy_limit_warps(kernel),  # noqa: F405
            get_sm__maximum_warps_avg_per_active_cycle(kernel),  # noqa: F405
            get_sm__maximum_warps_per_active_cycle_pct(kernel),  # noqa: F405
            get_sm__warps_active_avg_per_cycle_active(kernel),  # noqa: F405
            get_sm__warps_active_avg_pct_of_peak_sustained_active(kernel),  # noqa: F405
            get_l1tex__t_sector_hit_rate_pct(kernel),  # noqa: F405
            get_l1tex__t_sector_pipe_lsu_mem_global_op_ld_hit_rate_pct(kernel),  # noqa: F405
            get_lts__t_sector_hit_rate_pct(kernel),  # noqa: F405
            read_requests,
            write_requests,
        ]
        app_results.extend([
            read_requests + write_requests,
            read_transactions,
            write_transactions,
            read_transactions + write_transactions,
            0.0 if read_requests == 0.0 else read_transactions / read_requests,
            0.0 if write_requests == 0.0 else write_transactions / write_requests,
            l2_read,
            l2_write,
        ])
        app_results.extend([
            l2_read + l2_write,
            get_dram__sectors_read_sum(kernel) + get_dram__sectors_write_sum(kernel),  # noqa: F405
            atom_requests,
            red_requests,
        ])
        app_results.extend([
            atom_requests + red_requests,
            get_gpc__cycles_elapsed_max(kernel),  # noqa: F405
            get_sm__cycles_active_avg(kernel),  # noqa: F405
            get_smsp__inst_executed_sum(kernel),  # noqa: F405
            ipc,
            ipc * get_gpc__cycles_elapsed_avg_per_second(kernel) * 1e-6,  # noqa: F405
            get_gpu__time_duration_sum(kernel),  # noqa: F405
            l1_total_requests,
            get_lts__t_requests_srcunit_tex_sum(kernel),  # noqa: F405
            get_dram__bytes_sum_per_second(kernel),  # noqa: F405
            get_gpu__compute_memory_access_throughput_avg_pct_of_peak_sustained_elapsed(kernel),  # noqa: F405
            get_gpu__compute_memory_request_throughput_avg_pct_of_peak_sustained_elapsed(kernel),  # noqa: F405
            get_sm__memory_throughput_avg_pct_of_peak_sustained_elapsed(kernel),  # noqa: F405
            get_dram__cycles_elapsed_avg_per_second(kernel),  # noqa: F405
        ])
        rows.append(app_results)
        print(
            f"{report_file_path.name:>81s} knums: {len(kernels):4d} "
            f"kernel-{knum + 1:4d}"
        )
    return rows


def write_workbook(
    *,
    report_paths: Sequence[Path],
    output_path: Path,
    ncu_report: Any,
    max_kernels: int = 100,
) -> None:
    import openpyxl

    workbook = (
        openpyxl.load_workbook(output_path)
        if output_path.exists()
        else openpyxl.Workbook()
    )
    sheet_ncu = workbook.create_sheet("NCU", 1)
    sheet_ncu.append(_header_row())
    for report_path in report_paths:
        for row in _kernel_rows(
            report_file_path=report_path,
            ncu_report=ncu_report,
            max_kernels=max_kernels,
        ):
            sheet_ncu.append(row)
    workbook.save(output_path)


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    try:
        report_paths = _resolve_report_paths(args)
        ncu_report = load_ncu_report_module(args.ncu_report_path)
    except (FileNotFoundError, ImportError) as exc:
        print(f"read_ncu_report: {exc}", file=sys.stderr)
        return 2
    write_workbook(
        report_paths=report_paths,
        output_path=args.output,
        ncu_report=ncu_report,
        max_kernels=args.max_kernels,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
